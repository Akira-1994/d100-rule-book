#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""D100 規則書 — Phase 1 原始資料擷取器。

把 data/source/*.xlsx 以「純機械式」的方式轉成可進 git 的文字檔，
讓往後試算表更新時，git diff 能直接指出哪一張表的哪一列哪一欄改了。

本工具刻意不做任何資料清理、欄位對齊或語意判斷 —— 那些屬於 Phase 2。
唯一的職責是忠實且「可重現」地映射：同一份 xlsx 重跑任意次數，
產出的每一個位元組都必須完全相同。

擷取規則（變更任何一條都必須同步提升 TOOL_VERSION）
---------------------------------------------------
1.  不寫入任何時間戳記或機器相關資訊。
2.  數值：整數值的浮點數輸出為整數字串（Excel 一律以 double 儲存，
    「難度 1」還原成 1 比 1.0 更貼近作者輸入）；其餘浮點數以 repr 輸出。
3.  布林值輸出 TRUE / FALSE；日期時間輸出 ISO 8601。
4.  文字：先 rstrip 去除尾端空白與換行（尾端空白一律視為雜訊），
    但絕不 lstrip —— 散文型規則以行首空白表達階層，具有意義。
5.  跳脫順序：反斜線優先，接著 tab、CR、LF 依序轉為兩字元跳脫序列，
    確保每一個儲存格必定落在 TSV 的單一欄位內。
6.  去除每列尾端的全空欄位，以及整張表尾端的全空列。
7.  合併儲存格只有左上角帶值（openpyxl 行為），其餘為空；
    合併範圍本身另外記錄於 _merges.tsv，因此純版面調整也偵測得到。
8.  輸出一律為 UTF-8（無 BOM）、LF 換行。
9.  檔名以工作表名稱命名（不加序號），工作表順序記錄在 manifest，
    如此在中間插入新表時不會造成所有檔案改名的假差異。

用法
----
    python tools/extract.py                 # 擷取並寫入 data/raw
    python tools/extract.py --check         # 只檢查 data/raw 是否與來源同步
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.stderr.write("缺少相依套件，請先執行：pip install -r tools/requirements.txt\n")
    raise SystemExit(2)


# 擷取規則的版本。任何會改變輸出內容的邏輯調整都必須提升此版本，
# 以便在 diff 報告中區分「上游試算表改了」與「我們的擷取邏輯改了」。
TOOL_VERSION = "1.0.0"

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE_DIR = REPO_ROOT / "data" / "source"
DEFAULT_RAW_DIR = REPO_ROOT / "data" / "raw"

MANIFEST_NAME = "manifest.json"
MERGES_NAME = "_merges.tsv"
FORMULAS_NAME = "_formulas.tsv"

_ILLEGAL_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


def sanitize_filename(name: str) -> str:
    """把工作表名稱轉成安全的檔名（Windows 相容）。"""
    safe = _ILLEGAL_FILENAME_CHARS.sub("_", name)
    safe = safe.rstrip(" .")  # Windows 不允許檔名以空白或句點結尾
    return safe or "_unnamed"


def format_cell(value) -> str:
    """把單一儲存格的值轉成 TSV 欄位字串。規則見模組說明第 2–5 點。"""
    if value is None:
        return ""
    if value is True:
        return "TRUE"
    if value is False:
        return "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return repr(value)
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()

    text = str(value).rstrip()
    text = text.replace("\\", "\\\\")
    text = text.replace("\t", "\\t").replace("\r", "\\r").replace("\n", "\\n")
    return text


def extract_sheet_rows(worksheet) -> list[list[str]]:
    """讀出一張工作表的所有內容，並修剪尾端的全空欄與全空列。"""
    rows: list[list[str]] = []
    for raw_row in worksheet.iter_rows(min_row=1, min_col=1, values_only=True):
        cells = [format_cell(v) for v in raw_row]
        while cells and cells[-1] == "":
            cells.pop()
        rows.append(cells)
    while rows and not rows[-1]:
        rows.pop()
    return rows


def render_tsv(rows: list[list[str]]) -> str:
    return "".join("\t".join(row) + "\n" for row in rows)


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _formula_text(value):
    """辨識儲存格是否為公式，回傳公式字串（含等號）或 None。"""
    if isinstance(value, str):
        return value if value.startswith("=") else None
    text = getattr(value, "text", None)  # openpyxl ArrayFormula
    if isinstance(text, str) and text.startswith("="):
        return text
    return None


def extract(source: Path) -> dict:
    """讀取 xlsx，回傳 {相對檔名: 檔案內容} 的完整輸出集合。"""
    wb_values = openpyxl.load_workbook(source, data_only=True)
    wb_formulas = openpyxl.load_workbook(source, data_only=False)

    outputs: dict = {}
    sheet_entries: list = []
    merge_lines: list = ["工作表\t合併範圍\n"]
    formula_lines: list = ["工作表\t儲存格\t公式\n"]
    used_filenames: dict = {}

    for index, ws_values in enumerate(wb_values.worksheets):
        name = ws_values.title
        ws_formulas = wb_formulas[name]

        rows = extract_sheet_rows(ws_values)
        content = render_tsv(rows)

        filename = sanitize_filename(name) + ".tsv"
        if filename in used_filenames:
            # 兩張工作表名稱清洗後撞名，退回加上索引以保證唯一。
            filename = sanitize_filename(name) + "~" + str(index) + ".tsv"
        used_filenames[filename] = name
        outputs[filename] = content

        merges = sorted(
            ws_formulas.merged_cells.ranges,
            key=lambda r: (r.min_row, r.min_col, r.max_row, r.max_col),
        )
        for rng in merges:
            merge_lines.append(f"{name}\t{rng}\n")

        formula_count = 0
        cached_none_count = 0
        for row in ws_formulas.iter_rows(min_row=1, min_col=1):
            for cell in row:
                formula = _formula_text(cell.value)
                if formula is None:
                    continue
                formula_count += 1
                formula_lines.append(
                    f"{name}\t{cell.coordinate}\t{format_cell(formula)}\n"
                )
                if ws_values[cell.coordinate].value is None:
                    cached_none_count += 1

        sheet_entries.append(
            {
                "index": index,
                "name": name,
                "file": filename,
                "state": ws_values.sheet_state,
                "rows": len(rows),
                "cols": max((len(r) for r in rows), default=0),
                "sha256": sha256_text(content),
                "merged_ranges": len(merges),
                "formulas": formula_count,
                "formulas_without_cached_value": cached_none_count,
            }
        )

    outputs[MERGES_NAME] = "".join(merge_lines)
    outputs[FORMULAS_NAME] = "".join(formula_lines)

    manifest = {
        "tool_version": TOOL_VERSION,
        "openpyxl_version": openpyxl.__version__,
        "source": {
            "file": source.name,
            "bytes": source.stat().st_size,
            "sha256": sha256_file(source),
        },
        "totals": {
            "sheets": len(sheet_entries),
            "rows": sum(s["rows"] for s in sheet_entries),
            "merged_ranges": sum(s["merged_ranges"] for s in sheet_entries),
            "formulas": sum(s["formulas"] for s in sheet_entries),
        },
        "sheets": sheet_entries,
    }
    outputs[MANIFEST_NAME] = json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    return outputs


def resolve_source(source_arg) -> Path:
    if source_arg:
        path = Path(source_arg)
        if not path.is_file():
            raise SystemExit(f"找不到來源檔案：{path}")
        return path
    candidates = sorted(DEFAULT_SOURCE_DIR.glob("*.xlsx"))
    candidates = [p for p in candidates if not p.name.startswith("~$")]
    if not candidates:
        raise SystemExit(f"在 {DEFAULT_SOURCE_DIR} 找不到任何 .xlsx 檔案。")
    if len(candidates) > 1:
        names = "、".join(p.name for p in candidates)
        raise SystemExit(f"來源目錄有多個 xlsx（{names}），請以 --source 指定。")
    return candidates[0]


def write_outputs(outputs: dict, raw_dir: Path) -> list:
    """寫出所有檔案並刪除不再產生的舊檔，回傳實際被改動的檔名。"""
    raw_dir.mkdir(parents=True, exist_ok=True)
    changed: list = []

    for filename, content in sorted(outputs.items()):
        target = raw_dir / filename
        if target.exists() and target.read_text(encoding="utf-8") == content:
            continue
        with target.open("w", encoding="utf-8", newline="\n") as handle:
            handle.write(content)
        changed.append(filename)

    for existing in sorted(raw_dir.iterdir()):
        if existing.is_file() and existing.name not in outputs:
            existing.unlink()
            changed.append(f"（已刪除）{existing.name}")

    return changed


def check_outputs(outputs: dict, raw_dir: Path) -> list:
    """比對記憶體中的擷取結果與磁碟現況，回傳不同步的檔名清單。"""
    problems: list = []
    for filename, content in sorted(outputs.items()):
        target = raw_dir / filename
        if not target.exists():
            problems.append(f"缺少 {filename}")
        elif target.read_text(encoding="utf-8") != content:
            problems.append(f"內容不符 {filename}")
    if raw_dir.exists():
        for existing in sorted(raw_dir.iterdir()):
            if existing.is_file() and existing.name not in outputs:
                problems.append(f"多餘檔案 {existing.name}")
    return problems


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="把 D100 規則書試算表擷取成可進 git 的原始文字檔。"
    )
    parser.add_argument("--source", help="來源 xlsx 路徑（預設自動尋找 data/source）")
    parser.add_argument(
        "--raw-dir", default=str(DEFAULT_RAW_DIR), help="輸出目錄（預設 data/raw）"
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help="不寫檔，只檢查輸出目錄是否與來源同步；不同步時以狀態碼 1 結束",
    )
    args = parser.parse_args(argv)

    source = resolve_source(args.source)
    raw_dir = Path(args.raw_dir)
    outputs = extract(source)

    manifest = json.loads(outputs[MANIFEST_NAME])
    totals = manifest["totals"]

    if args.check:
        problems = check_outputs(outputs, raw_dir)
        if problems:
            print(f"data/raw 與 {source.name} 不同步：")
            for problem in problems:
                print(f"  - {problem}")
            print("請執行：python tools/extract.py")
            return 1
        print(f"data/raw 與 {source.name} 同步（{totals['sheets']} 張工作表）。")
        return 0

    changed = write_outputs(outputs, raw_dir)
    print(f"來源：{source.name}")
    print(f"雜湊：{manifest['source']['sha256']}")
    print(
        f"擷取：{totals['sheets']} 張工作表、{totals['rows']} 列、"
        f"{totals['merged_ranges']} 個合併範圍、{totals['formulas']} 條公式"
    )
    if changed:
        print(f"異動檔案 {len(changed)} 個：")
        for filename in changed:
            print(f"  - {filename}")
    else:
        print("輸出無變化。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
